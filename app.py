import streamlit as st
import pandas as pd
import io
import re
import sqlite3
import hashlib
import unicodedata
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from rapidfuzz import fuzz, process

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(page_title="PriceBot PRO V4.2", layout="wide")

# --- FUNÇÕES DE APOIO ---
def normalizar(txt):
    if not txt: return ""
    txt = str(txt).lower().strip()
    # Remove acentos para a similaridade funcionar bem
    txt = "".join(c for c in unicodedata.normalize('NFD', txt) if unicodedata.category(c) != 'Mn')
    return txt

def extrair_detalhes(texto):
    if not texto: return set()
    texto = normalizar(texto).replace(',', '.')
    # Padroniza medidas para evitar erro de 1kg vs 1 kg
    texto = re.sub(r'(\d+)\s+(g|kg|l|ml|lt|und|mts)', r'\1\2', texto)
    padrao = r'(\d+(?:\.\d+)?\s?(?:g|gr|kg|l|lt|ml|mts|und)\b)'
    return set(re.findall(padrao, texto))

# --- REGRA DE BARRAS ORIGINAL (COMO NO SEU 1º CÓDIGO) ---
def limpar_barcode_original(val):
    if pd.isna(val): return ""
    # Apenas remove o .0 e converte pra string (Exatamente como você fazia)
    return str(val).split('.')[0]

# --- BANCO DE DADOS ---
DB_NAME = "data_master.db"

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS products (description TEXT, barcode TEXT, price REAL)''')
    c.execute('''CREATE TABLE IF NOT EXISTS users (username TEXT PRIMARY KEY, password TEXT, expiry TEXT, role TEXT)''')
    if not c.execute("SELECT * FROM users WHERE username='admin'").fetchone():
        pw = hashlib.sha256("admin123".encode()).hexdigest()
        c.execute("INSERT INTO users VALUES (?, ?, ?, ?)", ('admin', pw, '2099-12-31', 'admin'))
    conn.commit()
    conn.close()

init_db()

# --- LOGIN ---
if 'auth' not in st.session_state: st.session_state.auth = False

if not st.session_state.auth:
    st.sidebar.title("🔐 Acesso")
    u, p = st.sidebar.text_input("Usuário"), st.sidebar.text_input("Senha", type="password")
    if st.sidebar.button("Entrar"):
        conn = sqlite3.connect(DB_NAME)
        pw = hashlib.sha256(p.encode()).hexdigest()
        res = conn.execute("SELECT role FROM users WHERE username=? AND password=?", (u, pw)).fetchone()
        if res:
            st.session_state.auth = True
            st.rerun()
        else: st.sidebar.error("Dados incorretos")
    st.stop()

# --- INTERFACE ---
menu = st.sidebar.radio("Navegação", ["📊 Processar Cotação", "⚙️ Gerenciar Banco"])

if menu == "📊 Processar Cotação":
    st.title("📊 Automatizador de Cotações PRO")

    if 'master_df' not in st.session_state:
        conn = sqlite3.connect(DB_NAME)
        st.session_state.master_df = pd.read_sql("SELECT * FROM products", conn)
        conn.close()

    master_df = st.session_state.master_df

    with st.sidebar:
        st.header("Configurações")
        modo_busca = st.selectbox("Escolha a regra:", 
                                 ["Híbrido (Barras + Similaridade)", 
                                  "Apenas Código de Barras", 
                                  "Apenas Similaridade"])
        sensibilidade = st.slider("Sensibilidade Similaridade (%)", 50, 100, 75)
        desconto = st.number_input("Desconto Global (%)", 0.0)
        debug = st.checkbox("🔍 Modo Debug")

    file = st.file_uploader("Upload da Cotação", type=["xlsx"])

    if file:
        h_row = st.number_input("Linha do Cabeçalho:", 1, 100, 10) # Padrão 10 como no seu original
        df_cols = pd.read_excel(file, header=h_row-1, nrows=0).columns.tolist()
        
        c1, c2, c3 = st.columns(3)
        col_desc = c1.selectbox("Coluna Descrição", df_cols)
        col_bar = c2.selectbox("Coluna Cód. Barras", df_cols)
        col_price = c3.selectbox("Coluna Preço", df_cols)

        if st.button("🚀 Iniciar Processamento"):
            # Mapeamento do banco
            price_map = dict(zip(master_df['barcode'].astype(str), master_df['price']))
            db_descs_norm = [normalizar(d) for d in master_df['description']]
            
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            header_map = {str(ws.cell(row=h_row, column=i).value).strip(): i for i in range(1, ws.max_column + 1)}
            idx_d, idx_b, idx_p = header_map[col_desc], header_map[col_bar], header_map[col_price]

            count, logs = 0, []

            for r in range(h_row + 1, ws.max_row + 1):
                orig_desc = str(ws.cell(row=r, column=idx_d).value or "")
                norm_desc = normalizar(orig_desc)
                # USA A REGRA ORIGINAL DE BARRAS
                orig_bar = limpar_barcode_original(ws.cell(row=r, column=idx_b).value)
                
                found_p = None
                status = "Não encontrado"

                # 1. BUSCA POR BARRAS (REGRA ORIGINAL)
                if "Barras" in modo_busca or "Híbrido" in modo_busca:
                    if orig_bar in price_map:
                        found_p = price_map[orig_bar]
                        status = "Match: Código de Barras"

                # 2. BUSCA POR SIMILARIDADE (REGRA NOVA TOP 5)
                if found_p is None and ("Similaridade" in modo_busca or "Híbrido" in modo_busca) and len(norm_desc) > 3:
                    alvo_pesos = extrair_detalhes(orig_desc)
                    matches = process.extract(norm_desc, db_descs_norm, scorer=fuzz.token_set_ratio, limit=5)
                    
                    for m_text, score, m_idx in matches:
                        if score >= sensibilidade:
                            item_db = master_df.iloc[m_idx]
                            db_pesos = extrair_detalhes(item_db['description'])
                            
                            if not alvo_pesos or not db_pesos or alvo_pesos == db_pesos:
                                found_p = item_db['price']
                                status = f"Match: {score}% similar"
                                break
                            else:
                                status = f"Bloqueado: Peso Divergente"

                if found_p:
                    v_final = float(found_p) * (1 - (desconto/100))
                    ws.cell(row=r, column=idx_p).value = round(v_final, 2)
                    count += 1
                
                if debug: logs.append({"Linha": r, "Produto": orig_desc, "Status": status})

            output = io.BytesIO()
            wb.save(output)
            st.success(f"Finalizado! {count} itens processados.")
            st.download_button("📥 Baixar Planilha", output.getvalue(), "cotacao_finalizada.xlsx")
            if debug: st.table(logs)

elif menu == "⚙️ Gerenciar Banco":
    st.title("⚙️ Gerenciar Banco de Dados")
    f_db = st.file_uploader("Upload Banco", type=["xlsx", "csv"])
    sub = st.checkbox("Substituir tudo?")
    if f_db and st.button("💾 Sincronizar"):
        df = pd.read_excel(f_db) if f_db.name.endswith('.xlsx') else pd.read_csv(f_db)
        df = df.iloc[:, [0, 1, 2]]
        df.columns = ['description', 'barcode', 'price']
        # TAMBÉM APLICA A REGRA ORIGINAL NA IMPORTAÇÃO
        df['barcode'] = df['barcode'].apply(limpar_barcode_original)
        df['price'] = pd.to_numeric(df['price'], errors='coerce').fillna(0.0)
        conn = sqlite3.connect(DB_NAME)
        df.to_sql("products", conn, if_exists="replace" if sub else "append", index=False)
        conn.close()
        if 'master_df' in st.session_state: del st.session_state['master_df']
        st.success("Banco Atualizado!")
